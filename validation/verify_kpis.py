"""Independent check of the Power BI KPIs.

Recalculates every headline measure directly from the raw Excel export
(no Power BI, no pandas) so the dashboard figures can be verified.

Usage:  pip install openpyxl
        python validation/verify_kpis.py
"""
import pathlib
PATH = pathlib.Path(__file__).resolve().parents[1] / "Data" / "MX-GOODS_Settlement_2026-10-01_to_2026-10-19.xlsx"
import openpyxl
from collections import defaultdict
from decimal import Decimal
wb=openpyxl.load_workbook(PATH,read_only=True,data_only=True)
ws=wb.worksheets[0]; it=ws.iter_rows(values_only=True); H=list(next(it)); ix={h:i for i,h in enumerate(H)}
def num(v):
    if v is None or str(v).strip()=='': return Decimal(0)
    return Decimal(str(v).replace(',','').strip())
rows=[r for r in it if any(c not in (None,'') for c in r)]
print('sheets',wb.sheetnames,'| data rows',len(rows),'| columns',len(H))
money=H[10:23]
print('currencies:',{r[ix['Currency']] for r in rows})
def cat(t):
    l=t.lower()
    if t=='Transfer': return 'Bank transfer'
    if 'withheld' in l and 'refund' in l: return 'Tax refunded'
    if l.startswith('marketplace withheld'): return 'Tax withheld'
    if 'shipping label' in l: return 'Shipping'
    if l.startswith('advertising'): return 'Advertising'
    if t=='Order Payment': return 'Sales'
    if t=='Refund': return 'Refunds'
    return 'Penalties & other fees'
C=defaultdict(Decimal); N=defaultdict(int); T=defaultdict(Decimal); TN=defaultdict(int)
paid=set(); active=set(); refunded=set(); units=0; gmv=Decimal(0); svc=Decimal(0)
st_pay=defaultdict(Decimal); st_ord=defaultdict(set); prod=defaultdict(Decimal); prodname={}
ostate={}; ocats=defaultdict(set)
for r in rows:
    t=r[ix['Transaction type']]; c=cat(t); tot=num(r[ix['Total']]); oid=(r[ix['Order ID']] or '').strip()
    C[c]+=tot; N[c]+=1; T[t]+=tot; TN[t]+=1
    if oid: active.add(oid); ocats[oid].add(c)
    st=(r[ix['Ship state']] or '').strip()
    if oid and st: ostate.setdefault(oid,st)
    if c=='Sales':
        paid.add(oid); units+=int(num(r[ix['Quantity']])); gmv+=num(r[ix['Retail price']]); svc+=num(r[ix['Service fee (tax incl.)']])
        st_pay[st]+=tot; st_ord[st].add(oid); sk=r[ix['SKU ID']]; prod[sk]+=tot; prodname[sk]=r[ix['SKU']]
    if c=='Refunds': refunded.add(oid)
net=sum(v for k,v in C.items() if k!='Bank transfer'); sales=C['Sales']
res={
 'Net Earnings':net,'Order Payments':sales,'Paid Orders':len(paid),'Units Sold':units,
 'Average Order Value':round(sales/len(paid),2),'Payout % of Sales':f"{net/sales*100:.1f}%",
 'Service Fees':svc,'Shipping Costs':C['Shipping'],'Advertising Spend':C['Advertising'],
 'Ad Spend % of Sales':f"{-C['Advertising']/sales*100:.1f}%",'Net Tax Impact':C['Tax withheld']+C['Tax refunded'],
 'Bank Transfers':C['Bank transfer'],'Refunds':C['Refunds'],'Tax Withheld':C['Tax withheld'],'Tax Refunded':C['Tax refunded'],
 'Penalties & Other Fees':C['Penalties & other fees'],'GMV':gmv,
 'Orders with Activity':len(active),'Orders without Payment':len([o for o in active if 'Sales' not in ocats[o]]),
 'Refund Rate %':f"{len(refunded)/len(paid)*100:.1f}% ({len(refunded)} orders)",'Shipping Cost per Order':round(-C['Shipping']/len(paid),2),
 'Ledger rows total':len(rows),'Balance movement (all rows)':sum(C.values())}
for k,v in res.items(): print(f"{k:32s} {v}")
print('\nMatrix rows:'); 
for c in ['Sales','Refunds','Shipping']: print(f"  {c:10s} {N[c]:4d} {C[c]}")
for t in ['Shipping label for return purchase','Shipping label for return purchase covered by platform','Shipping label purchase']: print(f"    {t:58s} {TN[t]:4d} {T[t]}")
print('\nTop states by order payments:'); 
for s,v in sorted(st_pay.items(),key=lambda x:-x[1])[:7]: print(f"  {s:34s} {v:>10}  orders {len(st_ord[s])}")
print('\nTop products:'); 
for k,v in sorted(prod.items(),key=lambda x:-x[1])[:4]: print(f"  {prodname[k][:45]:45s} {v}")
